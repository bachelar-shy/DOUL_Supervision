"""A4 derivative export from an already authorized, immutable HTML snapshot.

This module does not look up current domain records, select new template fields,
assert legal signatures, or replace received statutory originals.
"""
from __future__ import annotations

from datetime import datetime, timezone
import fcntl
from hashlib import sha256
from html.parser import HTMLParser
import io
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile
from typing import Protocol
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from app.reference.engine import canonical

VERSION = 'doul-a4-docx-pdf/1.0'
MIME = {'docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'pdf':'application/pdf'}


class DocumentExportPort(Protocol):
    def export(self, snapshot: dict, preserved_html: bytes, format: str) -> dict: ...


class FrozenHTML(HTMLParser):
    ALLOWED = {'html','meta','style','body','h1','p','b','table','thead','tbody','tr','th','td','div','br'}
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []
        self.style = False
        self.active = None
        self.parts = []
        self.table = None
        self.row = None
        self.heading = False
        self.title_count = 0

    def handle_decl(self, decl):
        if decl.casefold() != 'doctype html': raise ValueError('Unsupported document declaration')

    def handle_starttag(self, tag, attrs):
        if tag not in self.ALLOWED: raise ValueError('Unsupported frozen document element')
        attrs = dict(attrs)
        if any(k.startswith('on') or k in {'src','href','background'} for k in attrs):
            raise ValueError('External or executable document content is forbidden')
        if tag == 'style': self.style = True
        if tag in {'h1','p','td','th'}:
            self.active = tag; self.parts = []
        if tag == 'h1': self.title_count += 1
        if tag == 'table':
            if self.table is not None: raise ValueError('Nested tables are unsupported')
            self.table = []
        if tag == 'tr': self.row = []
        if tag == 'br' and self.active: self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag == 'style': self.style = False
        if tag in {'h1','p','th','td'}:
            value = ''.join(self.parts)
            if tag in {'h1','p'}: self.blocks.append((tag,value))
            else:
                if self.row is None: raise ValueError('Orphan table cell')
                self.row.append(value)
            self.active = None; self.parts = []
        if tag == 'tr':
            if self.table is None or self.row is None: raise ValueError('Orphan table row')
            self.table.append(self.row); self.row = None
        if tag == 'table':
            if not self.table or not self.table[0] or any(len(row) != len(self.table[0]) for row in self.table):
                raise ValueError('Inconsistent table columns')
            self.blocks.append(('table',self.table)); self.table = None

    def handle_data(self, value):
        if not self.style and self.active: self.parts.append(value)


def validate_source(snapshot, html):
    if snapshot.get('rendererVersion') not in {'reference-html/1', 'reference-html/2'} or not re.fullmatch(r'FM-[0-9]{2}', snapshot.get('formId','')):
        raise ValueError('Only reviewed internal-form snapshot HTML is supported')
    if not isinstance(html, bytes) or len(html) > 16 * 1024**2:
        raise ValueError('Invalid preserved document bytes')
    if sha256(html).hexdigest() != snapshot.get('renderHash'):
        raise ValueError('Preserved HTML hash mismatch')
    if sha256(canonical(snapshot['sourcePayload'])).hexdigest() != snapshot.get('sourcePayloadHash'):
        raise ValueError('Frozen source payload hash mismatch')
    if snapshot['rendererVersion'] == 'reference-html/2':
        catalog=snapshot['sourcePayload'].get('localization')
        if not isinstance(catalog,dict) or catalog.get('formId') != snapshot['formId'] or catalog.get('language') != snapshot['sourcePayload'].get('language'):
            raise ValueError('Localized renderer requires its exact frozen display catalog')
        catalog_body={key:value for key,value in catalog.items() if key != 'sha256'}
        if sha256(canonical(catalog_body)).hexdigest() != catalog.get('sha256'):
            raise ValueError('Frozen localization catalog hash mismatch')
    refs = [source['ref'] for source in snapshot['sourcePayload']['sources']]
    if refs != snapshot.get('sourceDomainRefs') or snapshot.get('primarySourceRef') not in refs:
        raise ValueError('Exact source references are inconsistent')
    if snapshot['sourcePayload']['language'] not in {'ko','vi','en'}:
        raise ValueError('Unsupported document language')
    if snapshot['formId'] in {'FM-09','FM-10','FM-11'} and snapshot['status'] != 'draft' and snapshot['sourcePayload']['language'] != 'vi':
        raise ValueError('Final statutory-language boundary must remain Vietnamese')
    parser = FrozenHTML(); parser.feed(html.decode('utf-8')); parser.close()
    if parser.title_count != 1 or parser.table is not None:
        raise ValueError('Invalid frozen document structure')
    return parser.blocks


def _wrap(value):
    # Soft breaks only: preserve the original values in the immutable source hash.
    return re.sub(r'\S{40,}', lambda m:'\u200b'.join(m.group()[i:i+24] for i in range(0,len(m.group()),24)), value)


def _fonts(run, size=10.5, bold=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name = 'Noto Sans'
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor(0,0,0)
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn('w:rFonts'))
    if fonts is None: fonts = OxmlElement('w:rFonts'); properties.append(fonts)
    for key, value in {'ascii':'Noto Sans','hAnsi':'Noto Sans','eastAsia':'Noto Sans CJK KR','cs':'Noto Sans'}.items(): fonts.set(qn('w:'+key), value)


def _no_split(row):
    from docx.oxml import OxmlElement
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))


def _cell_padding(cell):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    margins = OxmlElement('w:tcMar')
    for side in ('top','left','bottom','right'):
        element=OxmlElement('w:'+side); element.set(qn('w:w'),'90'); element.set(qn('w:type'),'dxa'); margins.append(element)
    cell._tc.get_or_add_tcPr().append(margins)


def make_docx(snapshot, html):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    blocks = validate_source(snapshot, html)
    document = Document(); section = document.sections[0]
    section.page_width=Mm(210); section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(17)
    section.header_distance=Mm(8); section.footer_distance=Mm(8)
    normal = document.styles['Normal']
    normal.font.name='Noto Sans'; normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.line_spacing=1.1
    language = snapshot['sourcePayload']['language']
    for kind, value in blocks:
        if kind in {'h1','p'}:
            paragraph=document.add_paragraph()
            paragraph.paragraph_format.keep_together = len(value) < 800
            if kind == 'h1':
                paragraph.paragraph_format.keep_with_next=True
                paragraph.paragraph_format.space_after=Pt(10)
                _fonts(paragraph.add_run(_wrap(value)), 17, True)
            else:
                label, separator, body = value.partition(': ')
                if separator:
                    _fonts(paragraph.add_run(_wrap(label)+': '),10.5,True)
                    _fonts(paragraph.add_run(_wrap(body)))
                else: _fonts(paragraph.add_run(_wrap(value)))
            continue
        rows = value; columns=len(rows[0])
        if columns > 16: raise ValueError('Too many columns for supported A4 form layout')
        weights=[]
        for index, header in enumerate(rows[0]):
            longest=max([len(str(row[index])) for row in rows[:30]] or [1])
            weights.append(max(8,min(32,max(len(header),longest)**0.6*4)))
        widths=[176*w/sum(weights) for w in weights]
        table=document.add_table(rows=0,cols=columns); table.style='Table Grid'; table.autofit=False
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for column,width in zip(table.columns,widths): column.width=Mm(width)
        for index, data in enumerate(rows):
            row=table.add_row()
            # Ordinary rows stay intact; exceptionally long rows may flow instead
            # of being clipped or pushed off the bottom of a page.
            if max(map(len,data)) < 600: _no_split(row)
            if index == 0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell,text,width in zip(row.cells,data,widths):
                cell.width=Mm(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _cell_padding(cell)
                paragraph=cell.paragraphs[0]; paragraph.paragraph_format.space_after=Pt(1)
                paragraph.paragraph_format.space_before=Pt(1)
                paragraph.paragraph_format.line_spacing=1.05
                if index == 0: paragraph.paragraph_format.keep_with_next=True
                if len(text)<18: paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                _fonts(paragraph.add_run(_wrap(text)), 9.5, index==0)
        document.add_paragraph().paragraph_format.space_after=Pt(3)
    roles=snapshot['sourcePayload'].get('requiredSignatureRoles',[])
    if roles:
        paragraph=document.add_paragraph()
        paragraph.paragraph_format.keep_with_next=True
        label={'ko':'서명란 (출력본)', 'vi':'Khu vực ký (bản xuất)', 'en':'Signature spaces (export copy)'}[language]
        _fonts(paragraph.add_run(label),11,True)
        table=document.add_table(rows=0,cols=2); table.style='Table Grid'; table.autofit=False
        for column in table.columns: column.width=Mm(88)
        for offset in range(0,len(roles),2):
            row=table.add_row(); row.height=Mm(23); row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST; _no_split(row)
            for index,cell in enumerate(row.cells):
                _cell_padding(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                role=roles[offset+index] if offset+index<len(roles) else ''
                _fonts(cell.paragraphs[0].add_run(_wrap(role)),9,True)
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    _fonts(footer.add_run('DOUL INTERNATIONAL | '+snapshot['formId']+' | '+str(snapshot['templateVersion'])+' | '),8)
    page=OxmlElement('w:fldSimple'); page.set(qn('w:instr'),'PAGE'); footer._p.append(page)
    paragraph=document.add_paragraph()
    paragraph.paragraph_format.space_before=Pt(5)
    note={'ko':'DOUL 내부 출력본. 원본 스냅샷에 대한 서명은 이 파생 파일의 전자서명을 의미하지 않습니다.',
          'vi':'Bản xuất nội bộ DOUL. Chữ ký của bản chụp nguồn không phải chữ ký điện tử của tệp xuất này.',
          'en':'DOUL internal export. A source snapshot signature does not digitally sign this derivative file.'}[language]
    _fonts(paragraph.add_run(note),8)
    properties=document.core_properties
    properties.title=next(value for kind,value in blocks if kind=='h1')
    properties.author='DOUL INTERNATIONAL'; properties.last_modified_by='DOUL document runtime'
    properties.subject='Source snapshot '+snapshot['id']; properties.language=language; properties.revision=1
    stamp=datetime.fromisoformat(snapshot['renderedAt'].replace('Z','+00:00')).astimezone(timezone.utc).replace(tzinfo=None)
    properties.created=stamp; properties.modified=stamp
    stream=io.BytesIO(); document.save(stream)
    # Stable ZIP metadata makes identical source/renderer inputs byte-reproducible.
    output=io.BytesIO()
    with ZipFile(stream) as source, ZipFile(output,'w',compression=ZIP_DEFLATED,compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info=ZipInfo(name,date_time=(1980,1,1,0,0,0)); info.compress_type=ZIP_DEFLATED
            info.external_attr=0o600<<16
            target.writestr(info,source.read(name))
    return output.getvalue()


class LocalDocumentRuntime:
    def __init__(self, cache, descriptor='/opt/doul-document-runtime.json', soffice='soffice'):
        self.cache=Path(cache); self.cache.mkdir(parents=True,exist_ok=True,mode=0o700)
        info=self.cache.lstat()
        import stat
        if not stat.S_ISDIR(info.st_mode) or info.st_mode & 0o077 or info.st_uid != os.getuid():
            raise ValueError('Document artifact cache must be private and owned by the renderer')
        self.descriptor=json.loads(Path(descriptor).read_text())
        if self.descriptor.get('rendererVersion') != VERSION or not self.descriptor.get('fonts'):
            raise ValueError('Installed renderer/font inventory is missing')
        for font in self.descriptor['fonts']:
            path=Path(font['path'])
            if not str(path).startswith('/usr/share/fonts/') or not path.is_file() or sha256(path.read_bytes()).hexdigest()!=font['sha256']:
                raise ValueError('Installed font inventory changed')
        self.fingerprint=sha256(canonical(self.descriptor)).hexdigest()
        self.soffice=soffice

    def export(self,snapshot,preserved_html,format):
        if format not in MIME: raise ValueError('Unsupported document export format')
        validate_source(snapshot,preserved_html)
        identity=sha256(canonical(dict(snapshotId=snapshot['id'],sourceRenderHash=snapshot['renderHash'],
                                      sourcePayloadHash=snapshot['sourcePayloadHash'],runtime=self.fingerprint))).hexdigest()
        directory=self.cache/identity
        lock_fd=os.open(self.cache/(identity+'.lock'),os.O_RDWR|os.O_CREAT|getattr(os,'O_NOFOLLOW',0),0o600)
        with os.fdopen(lock_fd,'a') as lock:
            fcntl.flock(lock.fileno(),fcntl.LOCK_EX)
            if not directory.exists():
                with tempfile.TemporaryDirectory(prefix='.document-',dir=self.cache) as temp:
                    work=Path(temp); docx=work/'artifact.docx'
                    docx.write_bytes(make_docx(snapshot,preserved_html)); docx.chmod(0o600)
                    home=work/'home'; home.mkdir(mode=0o700)
                    profile=work/'lo-profile'
                    env=dict(os.environ,HOME=str(home),TZ='UTC',LANG='C.UTF-8',LC_ALL='C.UTF-8',SAL_USE_VCLPLUGIN='svp')
                    result=subprocess.run([self.soffice,'-env:UserInstallation='+profile.as_uri(),
                        '--headless','--nologo','--nodefault','--norestore','--convert-to','pdf:writer_pdf_Export',
                        '--outdir',str(work),str(docx)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,env=env,timeout=90)
                    pdf=work/'artifact.pdf'
                    if result.returncode or not pdf.is_file() or not pdf.read_bytes().startswith(b'%PDF-'):
                        raise RuntimeError('Document PDF conversion failed')
                    pdf.chmod(0o600)
                    manifest=dict(schemaVersion=1,artifactId=identity,sourceSnapshotId=snapshot['id'],
                        sourceRenderHash=snapshot['renderHash'],sourcePayloadHash=snapshot['sourcePayloadHash'],
                        sourceDomainRefs=snapshot['sourceDomainRefs'],templateHash=snapshot['templateHash'],
                        templateVersion=snapshot['templateVersion'],rendererVersion=VERSION,runtimeFingerprint=self.fingerprint,
                        legalSignatureTransferred=False,outputs={kind:sha256((work/('artifact.'+kind)).read_bytes()).hexdigest() for kind in MIME})
                    # Only validated outputs are published. Profiles/font binaries
                    # never enter the output artifact or become downloadable.
                    publish=work/'publish'; publish.mkdir(mode=0o700)
                    for path in (docx,pdf): path.rename(publish/path.name)
                    (publish/'manifest.json').write_bytes(canonical(manifest)); (publish/'manifest.json').chmod(0o600)
                    publish.rename(directory)
            if directory.is_symlink(): raise ValueError('Artifact directory cannot be a symlink')
            if (directory/'manifest.json').is_symlink(): raise ValueError('Artifact manifest cannot be a symlink')
            manifest=json.loads((directory/'manifest.json').read_text())
            if manifest.get('artifactId')!=identity or manifest.get('sourceSnapshotId')!=snapshot['id'] or manifest.get('sourceRenderHash')!=snapshot['renderHash']:
                raise ValueError('Artifact source identity mismatch')
            path=directory/('artifact.'+format)
            if path.is_symlink(): raise ValueError('Artifact cannot be a symlink')
            content=path.read_bytes()
            if manifest['runtimeFingerprint'] != self.fingerprint or sha256(content).hexdigest()!=manifest['outputs'][format]:
                raise ValueError('Preserved document artifact integrity check failed')
            return dict(content=content,mediaType=MIME[format],sha256=manifest['outputs'][format],manifest=manifest)
